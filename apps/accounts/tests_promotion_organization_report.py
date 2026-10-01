from __future__ import annotations

from datetime import datetime, timedelta, timezone as datetime_timezone
from decimal import Decimal
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from openpyxl import load_workbook

from apps.accounts.models import (
    AlipayWebsitePayment,
    Entitlement,
    Module,
    ModuleSeason,
    PaymentDiscountApplication,
    PromotionCodeRecord,
    PurchaseOffer,
    UserCoupon,
)


class PromotionOrganizationExcelReportTests(TestCase):
    def setUp(self) -> None:
        self.user = get_user_model().objects.create_user(
            telephone="13600136000",
            country_code="+86",
            password="pass-123456",
        )
        learning_module, _ = Module.objects.get_or_create(
            key="learning_by_video",
            defaults={"name": "Learning by Video", "is_active": True},
        )
        exam_module, _ = Module.objects.get_or_create(
            key="exam_preparation",
            defaults={"name": "备考季", "is_active": True},
        )
        science_season = ModuleSeason.objects.create(
            module=learning_module,
            season_number=1,
            title="Science Season 1",
        )
        vlog_season = ModuleSeason.objects.create(
            module=learning_module,
            season_number=4,
            title="Vlog季",
        )
        self.science_offer = self._create_offer(
            "science-30d", "科普季 30 天", learning_module, science_season, Entitlement.Plan.MONTH_1
        )
        self.vlog_offer = self._create_offer(
            "vlog-lifetime", "Vlog季永久", learning_module, vlog_season, Entitlement.Plan.LIFETIME
        )
        self.exam_offer = self._create_offer(
            "exam-90d", "备考季 90 天", exam_module, None, Entitlement.Plan.MONTH_3
        )

    def _create_offer(self, code, title, module, season, plan) -> PurchaseOffer:
        return PurchaseOffer.objects.create(
            code=code,
            title=title,
            module=module,
            season=season,
            plan=plan,
            price_amount=Decimal("99.00"),
            currency="CNY",
        )

    def _create_applied_purchase(
        self,
        *,
        code: str,
        organization: str,
        offer: PurchaseOffer,
        applied_at: datetime,
        final_amount: str,
    ) -> PaymentDiscountApplication:
        promotion = PromotionCodeRecord.objects.create(
            code=code,
            campaign_name="渠道活动",
            organization_name=organization,
            remark="报表测试",
            discount_amount=Decimal("10.00"),
            minimum_order_amount=Decimal("0.00"),
            status=PromotionCodeRecord.Status.CONSUMED,
            consumed_by_user=self.user,
            consumed_at=applied_at - timedelta(minutes=1),
        )
        coupon = UserCoupon.objects.create(
            user=self.user,
            promotion_code=promotion,
            discount_amount=Decimal("10.00"),
            minimum_order_amount=Decimal("0.00"),
            status=UserCoupon.Status.USED,
        )
        payment = AlipayWebsitePayment.objects.create(
            merchant_order_no=f"ORDER-{code}",
            subject=offer.title,
            total_amount=Decimal(final_amount),
            status=AlipayWebsitePayment.Status.PAID,
            paid_at=applied_at,
        )
        coupon.used_payment = payment
        coupon.used_at = applied_at
        coupon.save(update_fields=["used_payment", "used_at", "updated_at"])
        return PaymentDiscountApplication.objects.create(
            payment=payment,
            coupon=coupon,
            promotion_code=promotion,
            user=self.user,
            offer=offer,
            original_amount=Decimal(final_amount) + Decimal("10.00"),
            automatic_discount_amount=Decimal("0.00"),
            promotion_discount_amount=Decimal("10.00"),
            final_amount=Decimal(final_amount),
            campaign_name_snapshot="渠道活动",
            campaign_organization_snapshot=organization,
            promotion_code_remark_snapshot="报表测试",
            status=PaymentDiscountApplication.Status.APPLIED,
            applied_at=applied_at,
        )

    def test_exports_microsecond_bounded_read_only_workbook_with_one_sheet_per_org(self) -> None:
        start = datetime(2026, 10, 1, 8, 0, 0, 123456, tzinfo=datetime_timezone.utc)
        end = datetime(2026, 10, 1, 18, 0, 0, 654321, tzinfo=datetime_timezone.utc)
        self._create_applied_purchase(
            code="SCIENCE001",
            organization="机构 A",
            offer=self.science_offer,
            applied_at=start,
            final_amount="49.90",
        )
        self._create_applied_purchase(
            code="VLOG000001",
            organization="机构 A",
            offer=self.vlog_offer,
            applied_at=end - timedelta(microseconds=1),
            final_amount="99.00",
        )
        self._create_applied_purchase(
            code="EXAM000001",
            organization="机构 B",
            offer=self.exam_offer,
            applied_at=start + timedelta(hours=1),
            final_amount="59.90",
        )
        self._create_applied_purchase(
            code="ATEND00001",
            organization="机构 A",
            offer=self.exam_offer,
            applied_at=end,
            final_amount="500.00",
        )
        before = list(
            PaymentDiscountApplication.objects.order_by("id").values_list(
                "id", "status", "applied_at", "final_amount"
            )
        )

        with TemporaryDirectory() as directory:
            output_path = Path(directory) / "promotion-report.xlsx"
            call_command(
                "promotion_organization_excel_report",
                start_at=start.isoformat(timespec="microseconds"),
                end_at=end.isoformat(timespec="microseconds"),
                output=str(output_path),
            )

            workbook = load_workbook(output_path, data_only=True)
            self.assertEqual(workbook.sheetnames, ["机构 A", "机构 B"])
            self.assertEqual(
                list(workbook["机构 A"].values),
                [
                    ("Promotion Code", "购买商品", "购买时长", "订单金额"),
                    ("SCIENCE001", "科普季", "30天", 49.9),
                    ("VLOG000001", "Vlog季", "永久", 99),
                    (None, None, "总计", 148.9),
                ],
            )
            self.assertEqual(
                list(workbook["机构 B"].values),
                [
                    ("Promotion Code", "购买商品", "购买时长", "订单金额"),
                    ("EXAM000001", "备考季", "90天", 59.9),
                    (None, None, "总计", 59.9),
                ],
            )

        self.assertEqual(
            list(
                PaymentDiscountApplication.objects.order_by("id").values_list(
                    "id", "status", "applied_at", "final_amount"
                )
            ),
            before,
        )
