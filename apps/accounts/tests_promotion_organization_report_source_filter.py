from datetime import datetime, timezone
from decimal import Decimal
from io import StringIO
from pathlib import Path
from tempfile import TemporaryDirectory

from django.contrib.auth import get_user_model
from django.core.management import call_command, CommandError
from django.test import TestCase
from openpyxl import load_workbook

from apps.accounts.models import (
    AlipayWebsitePayment,
    Entitlement,
    Module,
    PaymentDiscountApplication,
    PromotionCodeRecord,
    PurchaseOffer,
    UserCoupon,
)


class PromotionOrganizationReportSourceFilterTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            telephone="13600136001", country_code="+86", password="pass-123456"
        )
        module, _ = Module.objects.get_or_create(
            key="exam_preparation", defaults={"name": "备考季", "is_active": True}
        )
        self.offer = PurchaseOffer.objects.create(
            code="report-source-test", title="Report test", module=module,
            plan=Entitlement.Plan.MONTH_3, price_amount=Decimal("99.00"), currency="CNY",
        )
        self.start = datetime(2026, 10, 1, tzinfo=timezone.utc)
        self.end = datetime(2026, 10, 2, tzinfo=timezone.utc)

    def _purchase(self, code, source_kind, amount):
        applied_at = self.start.replace(hour=12)
        promotion = PromotionCodeRecord.objects.create(
            code=code, source_kind=source_kind, campaign_name="同名活动",
            organization_name="同名机构", remark="测试",
            discount_amount=Decimal("10.00"), minimum_order_amount=Decimal("0.00"),
            status=PromotionCodeRecord.Status.CONSUMED,
            consumed_by_user=self.user, consumed_at=applied_at,
        )
        coupon = UserCoupon.objects.create(
            user=self.user, promotion_code=promotion, discount_amount=Decimal("10.00"),
            minimum_order_amount=Decimal("0.00"), status=UserCoupon.Status.USED,
        )
        payment = AlipayWebsitePayment.objects.create(
            merchant_order_no=f"ORDER-{code}", subject=self.offer.title,
            total_amount=Decimal(amount), status=AlipayWebsitePayment.Status.PAID,
            paid_at=applied_at,
        )
        coupon.used_payment = payment
        coupon.used_at = applied_at
        coupon.save(update_fields=["used_payment", "used_at", "updated_at"])
        return PaymentDiscountApplication.objects.create(
            payment=payment, coupon=coupon, promotion_code=promotion, user=self.user,
            offer=self.offer, original_amount=Decimal(amount) + Decimal("10.00"),
            automatic_discount_amount=Decimal("0.00"), promotion_discount_amount=Decimal("10.00"),
            final_amount=Decimal(amount), campaign_name_snapshot="同名活动",
            campaign_organization_snapshot="同名机构", promotion_code_remark_snapshot="测试",
            status=PaymentDiscountApplication.Status.APPLIED, applied_at=applied_at,
        )

    def _run(self, output, stdout=None):
        return call_command(
            "promotion_organization_excel_report", start_at=self.start.isoformat(),
            end_at=self.end.isoformat(), output=str(output), stdout=stdout,
        )

    def test_workbook_and_success_summary_include_only_operational_source(self):
        self._purchase("OPSOURCE001", PromotionCodeRecord.SourceKind.OPERATIONAL, "49.90")
        self._purchase("REFSOURCE001", PromotionCodeRecord.SourceKind.REFERRAL, "29.90")
        self._purchase("INTSOURCE001", PromotionCodeRecord.SourceKind.INTERNAL_REWARD, "19.90")
        before = list(PaymentDiscountApplication.objects.values_list("id", "status", "final_amount"))
        with TemporaryDirectory() as directory:
            output = Path(directory) / "report.xlsx"
            stdout = StringIO()
            self._run(output, stdout=stdout)
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(workbook.sheetnames, ["同名活动"])
            self.assertEqual(list(workbook["同名活动"].values), [
                ("Promotion Code", "购买商品", "购买时长", "订单金额"),
                ("OPSOURCE001", "备考季", "90天", 49.9),
                (None, None, "总计", 49.9),
            ])
            self.assertIn("Exported 1 purchases across 1 campaigns", stdout.getvalue())
        self.assertEqual(
            list(PaymentDiscountApplication.objects.values_list("id", "status", "final_amount")),
            before,
        )

    def test_referral_and_internal_only_raise_no_data_without_output(self):
        self._purchase("REFSOURCE002", PromotionCodeRecord.SourceKind.REFERRAL, "29.90")
        self._purchase("INTSOURCE002", PromotionCodeRecord.SourceKind.INTERNAL_REWARD, "19.90")
        with TemporaryDirectory() as directory:
            output = Path(directory) / "report.xlsx"
            with self.assertRaisesMessage(CommandError, "No paid promotion-code purchases"):
                self._run(output)
            self.assertFalse(output.exists())
